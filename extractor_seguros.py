"""
Extractor de Resúmenes de Cuenta de Pólizas de Seguro PDF → Excel
Procesa PDFs con formato "RESUMEN DE DEUDA" de aseguradoras (ej: Zurich,荔, etc.)
"""

import re
import pdfplumber
from datetime import datetime
from typing import Optional


ROW_PATTERN = re.compile(
    r'(\d{1,3}(?:\.\d{3})*,\d{3}/\d{3})'
    r'(\d{2}/\d{2}/\d{4})\s+'
    r'(\d{2}/\d{2}/\d{4})\s+\$\s+'
    r'([\d\.]+,\d{2})\s+'
    r'([A-Z])\s+'
    r'(\d{2}/\d{2}/\d{4})\s+'
    r'([\d\.,]+)\s+'
    r'([\d\.,]+)'
    r'(.*)'
)

HEADER_PATTERNS = {
    'fecha_resumen': re.compile(r'RESUMEN DE DEUDA AL\s+(.+)'),
    'cliente': re.compile(r'CLIENTE\s+(.+)'),
    'domicilio': re.compile(r'DOMICILIO\s+(.+)'),
    'telefono': re.compile(r'TELEFONO\s+(.+)'),
    'contacto_estado': re.compile(r'CONTACTO\s+(.+)'),
    'productor': re.compile(r'PRODUCTOR:\s+(.+)'),
}


def extraer_resumen_seguros(ruta_pdf: str) -> dict:
    """
    Extrae todas las filas de un PDF de resumen de deuda de pólizas de seguro.

    Args:
        ruta_pdf: Ruta al archivo PDF

    Returns:
        dict con:
          - 'polizas': lista de dicts con datos de cada póliza
          - 'info': metadata del resumen (fecha, cliente, productor, etc.)
          - 'total_polizas': cantidad de pólizas encontradas
    """
    polizas = []
    info = {}
    pagina_actual = 0

    with pdfplumber.open(ruta_pdf) as pdf:
        total_paginas = len(pdf.pages)

        for page in pdf.pages:
            pagina_actual += 1
            text = page.extract_text()
            if not text:
                continue

            for line in text.split('\n'):
                line = line.strip()
                if not line:
                    continue

                # Capturar header informativo
                for key, pattern in HEADER_PATTERNS.items():
                    m = pattern.match(line)
                    if m:
                        info[key] = m.group(1).strip()

                # Capturar filas de datos
                m = ROW_PATTERN.match(line)
                if m:
                    polizas.append({
                        'poliza': m.group(1),
                        'vigencia_desde': m.group(2),
                        'vigencia_hasta': m.group(3),
                        'saldo': m.group(4),
                        'tp': m.group(5),
                        'vencimiento': m.group(6),
                        'interes': m.group(7),
                        'factura': m.group(8),
                        'asegurado_objeto': m.group(9).strip(),
                    })

    info['paginas'] = pagina_actual
    return {
        'polizas': polizas,
        'info': info,
        'total_polizas': len(polizas),
    }


def exportar_excel(resultado: dict, ruta_excel: str) -> str:
    """
    Exporta el resultado a un archivo Excel formateado.

    Args:
        resultado: dict devuelto por extraer_resumen_seguros()
        ruta_excel: Ruta donde guardar el Excel

    Returns:
        Ruta del archivo generado
    """
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    polizas = resultado['polizas']
    info = resultado['info']

    wb = Workbook()
    ws = wb.active
    ws.title = "Polizas"

    columnas = [
        'POLIZA', 'VIGENCIA_DESDE', 'VIGENCIA_HASTA', 'SALDO', 'TP',
        'VENCIMIENTO', 'INTERES', 'FACTURA', 'ASEGURADO_OBJETO'
    ]

    h_font = Font(bold=True, color="FFFFFF", size=11)
    h_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side('thin'), right=Side('thin'),
        top=Side('thin'), bottom=Side('thin')
    )

    for ci, h in enumerate(columnas, 1):
        c = ws.cell(1, ci, h)
        c.font = h_font
        c.fill = h_fill
        c.alignment = h_align
        c.border = border

    for ri, row in enumerate(polizas, 2):
        for ci, col in enumerate(columnas, 1):
            c = ws.cell(ri, ci, row.get(col.lower(), ''))
            c.border = border
            c.alignment = Alignment(vertical="center")

    anchos = [16, 16, 16, 14, 5, 16, 10, 14, 60]
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(polizas)+1}"

    # Hoja Resumen
    ws2 = wb.create_sheet("Resumen", 0)
    resumen_data = [
        ('Campo', 'Valor'),
        ('Fecha del Resumen', info.get('fecha_resumen', '')),
        ('Cliente', limpiar_header(info.get('cliente', ''))),
        ('Domicilio', limpiar_header(info.get('domicilio', ''))),
        ('Productor', limpiar_header(info.get('productor', ''))),
        ('Total Pólizas', resultado['total_polizas']),
        ('Total Páginas', info.get('paginas', 0)),
        ('Fecha Exportación', datetime.now().strftime('%d/%m/%Y %H:%M')),
    ]
    for ri, (k, v) in enumerate(resumen_data, 1):
        ck = ws2.cell(ri, 1, k)
        cv = ws2.cell(ri, 2, v)
        if ri == 1:
            ck.font = h_font
            ck.fill = h_fill
            cv.font = h_font
            cv.fill = h_fill
        ck.border = border
        cv.border = border

    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 60

    wb.save(ruta_excel)
    return ruta_excel


def limpiar_header(texto: str) -> str:
    """Limpia artefactos numéricos del header (ej: 'CLIENTE 1 5 ,364' → 'CLIENTE 15,364')"""
    texto = re.sub(r'(\d)\s+(\d)', r'\1\2', texto)
    texto = re.sub(r'(\d)\s+(?=,)', r'\1', texto)
    return texto


def limpiar_texto_asegurado(texto: str) -> str:
    """
    Limpia caracteres numéricos espurios que aparecen en el texto extraído del PDF.
    Ej: 'bel1e5n49' → 'belen', 'CECILIA1549' → 'CECILIA'
    """
    if not texto:
        return texto
    texto = re.sub(r'(?<=[a-zA-Z])[0-9](?=[a-zA-Z])', '', texto)
    texto = re.sub(r'(?<=[a-zA-Z])[0-9]{2,4}(?=[a-zA-Z\s]|$)', '', texto)
    texto = re.sub(r'(?<=[a-zA-Z])[0-9](?=\s|$)', '', texto)
    texto = re.sub(r'\s+', ' ', texto).strip()
    return texto


def procesar_pdf(ruta_pdf: str, ruta_excel: Optional[str] = None) -> str:
    """
    Función de alto nivel: extrae y exporta en un solo paso.

    Args:
        ruta_pdf: Ruta al PDF de resumen de seguro
        ruta_excel: Ruta de salida para el Excel (auto si None)

    Returns:
        Ruta del Excel generado
    """
    if ruta_excel is None:
        base = ruta_pdf.rsplit('.', 1)[0]
        ruta_excel = f"{base}.xlsx"

    resultado = extraer_resumen_seguros(ruta_pdf)
    exportar_excel(resultado, ruta_excel)
    return ruta_excel
